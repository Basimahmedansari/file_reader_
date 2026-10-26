import json
import pandas as pd
from pypdf import PdfReader
from docx import Document
from pptx import Presentation


def read_file(file):
    name = file.name.lower()
    file.seek(0)

    if name.endswith(".pdf"):
        reader = PdfReader(file)
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    elif name.endswith(".docx"):
        doc = Document(file)
        return "\n".join(
            paragraph.text
            for paragraph in doc.paragraphs
        )

    elif name.endswith((".txt", ".md")):
        return file.getvalue().decode(
            "utf-8", errors="replace"
        )

    elif name.endswith(".csv"):
        df = pd.read_csv(file)
        return df.to_string(index=False)

    elif name.endswith(".xlsx"):
        sheets = pd.read_excel(
            file, sheet_name=None
        )
        return "\n\n".join(
            f"Sheet: {sheet}\n{df.to_string(index=False)}"
            for sheet, df in sheets.items()
        )

    elif name.endswith(".json"):
        data = json.loads(
            file.getvalue().decode("utf-8")
        )
        return json.dumps(
            data, ensure_ascii=False, indent=2
        )

    elif name.endswith(".pptx"):
        presentation = Presentation(file)
        slides = []

        for number, slide in enumerate(
            presentation.slides, start=1
        ):
            content = [
                shape.text
                for shape in slide.shapes
                if shape.has_text_frame
                and shape.text.strip()
            ]
            slides.append(
                f"Slide {number}:\n"
                + "\n".join(content)
            )

        return "\n\n".join(slides)

    else:
        raise ValueError(
            f"Unsupported file type: {name}"
        )